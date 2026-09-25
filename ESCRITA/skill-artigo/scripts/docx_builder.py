#!/usr/bin/env python3
"""
docx_builder.py — monta o .docx final do artigo científico. Aplica formatação
mecânica fixa (fonte, margens, espaçamento, blocos de resumo/abstract, numeração
opcional de seção) — nunca gera conteúdo, só monta o que já foi redigido.

Uso:
  python docx_builder.py --config artigo.json --out artigo.docx

Formato de --config (ver assets/config_exemplo.json):
{
  "meta": {
    "titulo": "...",
    "autores": ["Nome Completo (afiliação/instituição)"],
    "resumo": "...", "palavras_chave": ["...", "...", "..."],
    "abstract": "...", "keywords": ["...", "...", "..."],
    "periodico_alvo": "...", "norma_citacao": "abnt|apa|vancouver",
    "nota_origem": ""   // opcional — ver aviso de autoplágio no SKILL.md
  },
  "secoes": [
    {"titulo": "Introdução", "texto": "Parágrafo 1.\n\nParágrafo 2."},
    {"titulo": "Método", "texto": "..."}
  ],
  "numerar_secoes": false,          // opcional, default false (comum em artigos)
  "referencias": ["..."],           // opcional (lista pronta)
  "referencias_log": "research_log.json"   // opcional (gera via reference_formatter)
}
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).parent))
from reference_formatter import gerar_lista_referencias  # noqa: E402
from research_log import carregar as carregar_log  # noqa: E402

FONTE = "Times New Roman"
TAM_TEXTO = Pt(12)
PRETO = RGBColor(0, 0, 0)


def _set_margens_e_fonte(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.left_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = TAM_TEXTO
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)


def _add_page_number_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def _paragrafo(doc, texto, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=True,
               tamanho=TAM_TEXTO, negrito=False, italico=False, espacamento=1.5):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    if recuo:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = espacamento
    run = p.add_run(texto)
    run.font.name = FONTE
    run.font.size = tamanho
    run.font.color.rgb = PRETO
    run.bold = negrito
    run.italic = italico
    return p


def _add_cabecalho(doc: Document, meta: dict) -> None:
    _paragrafo(doc, meta.get("titulo", "[TÍTULO A PREENCHER]"),
               alinhamento=WD_ALIGN_PARAGRAPH.CENTER, recuo=False, tamanho=Pt(14), negrito=True)
    doc.add_paragraph()

    autores = meta.get("autores", [])
    if autores:
        _paragrafo(doc, "; ".join(autores), alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                   recuo=False, tamanho=Pt(11))
    doc.add_paragraph()

    if meta.get("nota_origem"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(meta["nota_origem"])
        run.font.name = FONTE
        run.font.size = Pt(9)
        run.italic = True
        doc.add_paragraph()


def _add_resumo_abstract(doc: Document, meta: dict) -> None:
    def bloco(rotulo, texto, termos, rotulo_termos):
        p_tit = doc.add_paragraph()
        run = p_tit.add_run(rotulo)
        run.font.name = FONTE; run.font.size = TAM_TEXTO; run.bold = True; run.font.color.rgb = PRETO
        _paragrafo(doc, texto, recuo=False, espacamento=1.0)
        if termos:
            p = doc.add_paragraph()
            run_lbl = p.add_run(f"{rotulo_termos}: ")
            run_lbl.font.name = FONTE; run_lbl.font.size = TAM_TEXTO; run_lbl.bold = True
            run_lbl.font.color.rgb = PRETO
            run_txt = p.add_run("; ".join(termos) + ".")
            run_txt.font.name = FONTE; run_txt.font.size = TAM_TEXTO; run_txt.font.color.rgb = PRETO
        doc.add_paragraph()

    bloco("RESUMO", meta.get("resumo", "[RESUMO A PREENCHER]"),
          meta.get("palavras_chave", []), "Palavras-chave")
    bloco("ABSTRACT", meta.get("abstract", "[ABSTRACT A PREENCHER]"),
          meta.get("keywords", []), "Keywords")


def _add_secoes(doc: Document, secoes: list, numerar: bool) -> None:
    contador = 0
    for secao in secoes:
        titulo = secao["titulo"].strip()
        texto = secao.get("texto", "")
        contador += 1
        prefixo = f"{contador} " if numerar else ""

        heading = doc.add_paragraph(style="Heading 1")
        heading.paragraph_format.line_spacing = 1.5
        run = heading.add_run(f"{prefixo}{titulo.upper()}")
        run.font.name = FONTE
        run.font.size = TAM_TEXTO
        run.font.color.rgb = PRETO
        run.bold = True

        for paragrafo in [p for p in texto.split("\n\n") if p.strip()]:
            _paragrafo(doc, paragrafo.strip())


def _add_referencias(doc: Document, referencias: list) -> None:
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = titulo.add_run("REFERÊNCIAS")
    run.font.name = FONTE; run.font.size = TAM_TEXTO; run.bold = True; run.font.color.rgb = PRETO
    doc.add_paragraph()
    for ref in referencias:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(ref)
        run.font.name = FONTE; run.font.size = TAM_TEXTO; run.font.color.rgb = PRETO


def montar_documento(config: dict) -> Document:
    doc = Document()
    _set_margens_e_fonte(doc)
    _add_page_number_footer(doc)

    meta = config.get("meta", {})
    _add_cabecalho(doc, meta)
    _add_resumo_abstract(doc, meta)
    _add_secoes(doc, config.get("secoes", []), config.get("numerar_secoes", False))

    referencias = config.get("referencias")
    estilo = meta.get("norma_citacao", "abnt")
    if not referencias and config.get("referencias_log"):
        log = carregar_log(Path(config["referencias_log"]))
        referencias = gerar_lista_referencias(log, estilo)
    if referencias:
        doc.add_paragraph()
        _add_referencias(doc, referencias)

    return doc


def _cli():
    ap = argparse.ArgumentParser(description="Monta o .docx final do artigo científico")
    ap.add_argument("--config", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    args = ap.parse_args()

    config = json.loads(args.config.read_text(encoding="utf-8"))
    doc = montar_documento(config)
    doc.save(str(args.out))
    print(f"Documento gerado em {args.out}")


if __name__ == "__main__":
    _cli()
