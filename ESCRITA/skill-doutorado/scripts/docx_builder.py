#!/usr/bin/env python3
"""
docx_builder.py — monta o .docx final aplicando as regras físicas fixas da
NBR 14724/15287 (fonte, margens, espaçamento, capa, folha de rosto, sumário,
numeração progressiva de seções, paginação). Recebe o TEXTO já redigido pelo
LLM por seção — nunca gera conteúdo, só aplica formatação mecânica.

Uso:
  python docx_builder.py --config projeto.json --out saida.docx

Formato de --config (ver assets/config_exemplo.json):
{
  "meta": {
    "instituicao": "...", "ppg": "...", "linha_pesquisa": "...",
    "autor": "...", "orientador": "...", "coorientador": "",
    "titulo": "...", "natureza": "Projeto de pesquisa apresentado ao ...",
    "local": "...", "ano": "2026"
  },
  "secoes": [
    {"nivel": 1, "titulo": "INTRODUÇÃO", "texto": "Parágrafo 1.\n\nParágrafo 2."},
    {"nivel": 2, "titulo": "Justificativa acadêmica", "texto": "..."}
  ],
  "cronograma": { ... saída de cronograma_builder.py ... },   // opcional
  "referencias": ["SILVA, João. Título. Veículo, 2023."],      // opcional (lista pronta)
  "referencias_log": "research_log.json"                       // opcional (gera via reference_formatter)
}
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PRETO = RGBColor(0, 0, 0)

sys.path.insert(0, str(Path(__file__).parent))
from reference_formatter import gerar_lista_referencias  # noqa: E402
from research_log import carregar as carregar_log  # noqa: E402

FONTE = "Times New Roman"
TAM_TEXTO = Pt(12)


# ---------------------------------------------------------------- infra ----

def _set_margens_e_fonte(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.left_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = TAM_TEXTO
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    # garante a fonte também para o script do East Asian (evita fallback)
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONTE)


def _add_page_number_footer(doc: Document) -> None:
    """Insere campo PAGE no rodapé, alinhado à direita (convenção: nº no canto
    superior direito; como python-docx não expõe cabeçalho com facilidade para
    todos os viewers, usamos rodapé direita — ajuste manual no Word se o padrão
    do programa exigir estritamente o canto superior)."""
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def _paragrafo_padrao(doc: Document, texto: str, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY,
                       recuo_primeira_linha=True):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    if recuo_primeira_linha:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(texto)
    run.font.name = FONTE
    run.font.size = TAM_TEXTO
    return p


# ------------------------------------------------------------- capa/rosto ----

def _add_capa(doc: Document, meta: dict) -> None:
    def centralizado(texto, tamanho=12, negrito=False, espaco_antes=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(espaco_antes)
        run = p.add_run(texto)
        run.font.name = FONTE
        run.font.size = Pt(tamanho)
        run.bold = negrito
        return p

    centralizado(meta.get("instituicao", "[INSTITUIÇÃO A PREENCHER]"), 12)
    centralizado(meta.get("ppg", "[PROGRAMA DE PÓS-GRADUAÇÃO A PREENCHER]"), 12)
    for _ in range(8):
        doc.add_paragraph()
    centralizado(meta.get("autor", "[AUTOR A PREENCHER]"), 12, negrito=True)
    for _ in range(4):
        doc.add_paragraph()
    centralizado(meta.get("titulo", "[TÍTULO PROVISÓRIO A PREENCHER]"), 14, negrito=True)
    for _ in range(10):
        doc.add_paragraph()
    centralizado(meta.get("local", "[LOCAL]"), 12)
    centralizado(str(meta.get("ano", "[ANO]")), 12)
    doc.add_page_break()


def _add_folha_rosto(doc: Document, meta: dict) -> None:
    def centralizado(texto, tamanho=12, negrito=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(texto)
        run.font.name = FONTE
        run.font.size = Pt(tamanho)
        run.bold = negrito
        return p

    centralizado(meta.get("autor", "[AUTOR A PREENCHER]"), 12, negrito=True)
    for _ in range(4):
        doc.add_paragraph()
    centralizado(meta.get("titulo", "[TÍTULO PROVISÓRIO A PREENCHER]"), 14, negrito=True)
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(8)
    run = p.add_run(meta.get("natureza", "[NATUREZA DO DOCUMENTO A PREENCHER]"))
    run.font.name = FONTE
    run.font.size = TAM_TEXTO

    if meta.get("linha_pesquisa"):
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(8)
        r2 = p2.add_run(f"Linha de pesquisa: {meta['linha_pesquisa']}")
        r2.font.name = FONTE
        r2.font.size = TAM_TEXTO

    p3 = doc.add_paragraph()
    p3.paragraph_format.left_indent = Cm(8)
    r3 = p3.add_run(f"Orientador(a): {meta.get('orientador', '[A PREENCHER]')}")
    r3.font.name = FONTE
    r3.font.size = TAM_TEXTO

    if meta.get("coorientador"):
        p4 = doc.add_paragraph()
        p4.paragraph_format.left_indent = Cm(8)
        r4 = p4.add_run(f"Coorientador(a): {meta['coorientador']}")
        r4.font.name = FONTE
        r4.font.size = TAM_TEXTO

    for _ in range(8):
        doc.add_paragraph()
    centralizado(meta.get("local", "[LOCAL]"), 12)
    centralizado(str(meta.get("ano", "[ANO]")), 12)
    doc.add_page_break()


def _add_sumario(doc: Document) -> None:
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("SUMÁRIO")
    run.font.name = FONTE
    run.font.size = TAM_TEXTO
    run.bold = True
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

    nota = doc.add_paragraph()
    nota_run = nota.add_run(
        "[Sumário gerado por campo — no Word, clique com o botão direito sobre "
        "esta área e escolha \"Atualizar campo\" após finalizar o texto.]"
    )
    nota_run.italic = True
    nota_run.font.size = Pt(9)
    doc.add_page_break()


# --------------------------------------------------------------- seções ----

def _add_secoes(doc: Document, secoes: list) -> None:
    contadores = [0, 0, 0]  # nível 1, 2, 3

    for secao in secoes:
        nivel = secao.get("nivel", 1)
        titulo = secao["titulo"].strip()
        texto = secao.get("texto", "")

        if nivel == 1:
            contadores[0] += 1
            contadores[1] = 0
            contadores[2] = 0
            numero = f"{contadores[0]}"
            doc.add_page_break()
        elif nivel == 2:
            contadores[1] += 1
            contadores[2] = 0
            numero = f"{contadores[0]}.{contadores[1]}"
        else:
            contadores[2] += 1
            numero = f"{contadores[0]}.{contadores[1]}.{contadores[2]}"

        estilo = f"Heading {min(nivel, 3)}"
        heading = doc.add_paragraph(style=estilo)
        heading.paragraph_format.line_spacing = 1.5
        run = heading.add_run(f"{numero} {titulo.upper() if nivel == 1 else titulo}")
        run.font.name = FONTE
        run.font.size = TAM_TEXTO
        run.font.color.rgb = PRETO  # ABNT exige preto; o estilo Heading padrão herda azul do tema
        run.bold = True
        if nivel == 3:
            run.italic = True
            run.bold = False

        for paragrafo in [p for p in texto.split("\n\n") if p.strip()]:
            _paragrafo_padrao(doc, paragrafo.strip())


# ----------------------------------------------------------- cronograma ----

def _add_cronograma(doc: Document, cronograma: dict) -> None:
    n_periodos = len(cronograma["periodos"])
    n_col = 1 + n_periodos
    table = doc.add_table(rows=1, cols=n_col)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # coluna "Etapa" larga, colunas de período estreitas e uniformes
    largura_util = Cm(21) - Cm(3) - Cm(2)  # A4 menos margens esq(3)+dir(2)
    largura_etapa = Cm(4.5)
    largura_periodo = (largura_util - largura_etapa) / n_periodos

    def _set_col_width(col_idx, largura):
        for row in table.rows:
            row.cells[col_idx].width = largura

    hdr = table.rows[0].cells
    hdr[0].text = "Etapa"
    for i, periodo in enumerate(cronograma["periodos"], start=1):
        hdr[i].text = periodo
    for linha in cronograma["linhas"]:
        row = table.add_row().cells
        row[0].text = linha["etapa"]
        for i, marca in enumerate(linha["marcacoes"], start=1):
            row[i].text = marca

    _set_col_width(0, largura_etapa)
    for i in range(1, n_col):
        _set_col_width(i, largura_periodo)

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = FONTE
                    run.font.size = Pt(10)
    table.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph()


# ----------------------------------------------------------- referências ----

def _add_referencias(doc: Document, referencias: list) -> None:
    doc.add_page_break()
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("REFERÊNCIAS")
    run.font.name = FONTE
    run.font.size = TAM_TEXTO
    run.bold = True
    doc.add_paragraph()

    for ref in referencias:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(ref)
        run.font.name = FONTE
        run.font.size = TAM_TEXTO


# --------------------------------------------------------------- build ----

def montar_documento(config: dict) -> Document:
    doc = Document()
    _set_margens_e_fonte(doc)
    _add_page_number_footer(doc)

    _add_capa(doc, config.get("meta", {}))
    _add_folha_rosto(doc, config.get("meta", {}))
    _add_sumario(doc)
    _add_secoes(doc, config.get("secoes", []))

    if config.get("cronograma"):
        doc.add_page_break()
        t = doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = t.add_run("CRONOGRAMA")
        r.font.name = FONTE
        r.font.size = TAM_TEXTO
        r.bold = True
        doc.add_paragraph()
        _add_cronograma(doc, config["cronograma"])

    referencias = config.get("referencias")
    if not referencias and config.get("referencias_log"):
        log = carregar_log(Path(config["referencias_log"]))
        referencias = gerar_lista_referencias(log)
    if referencias:
        _add_referencias(doc, referencias)

    return doc


def _cli():
    ap = argparse.ArgumentParser(description="Monta o .docx final conforme NBR 14724/15287")
    ap.add_argument("--config", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    args = ap.parse_args()

    config = json.loads(args.config.read_text(encoding="utf-8"))
    doc = montar_documento(config)
    doc.save(str(args.out))
    print(f"Documento gerado em {args.out}")
    print("Lembrete: abra no Word/LibreOffice e atualize o campo do sumário "
          "(botão direito → Atualizar campo) antes de considerar final.")


if __name__ == "__main__":
    _cli()
